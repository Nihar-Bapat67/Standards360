"""Module B1: Input Handler.

Accepts all four input types the manual lists, a tender PDF, pasted text, a screenshot, or a bare
product name, and reduces every one of them to plain text plus a note of where it came from. It also
records how rich the input is, because that single field decides how hard B5 has to work later.

Screenshots need OCR. Tesseract is optional here: when it is missing, B1 says so in `notes` rather
than failing, because the rest of the pipeline still works for the other three input types.

    from app.understand.input_handler import InputHandler
    InputHandler().read("tender.pdf")
    InputHandler().read_text("43 grade OPC for RCC work")
"""

import re
import sys
from pathlib import Path
from typing import List, Optional

sys.path.insert(0, str(Path(__file__).resolve().parent.parent.parent))

from contracts.requirement import InputPayload  # noqa: E402

PDF_SUFFIXES = {".pdf"}
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".webp"}
DOC_SUFFIXES = {".docx"}
TEXT_SUFFIXES = {".txt", ".md"}

# Headings that mark the start of the part of a tender we care about.
SPEC_HEADING = re.compile(
    r"\b(technical\s+specification|specification\s+of\s+(?:works|materials|goods)|scope\s+of\s+(?:work|supply)|"
    r"schedule\s+of\s+(?:quantities|items)|bill\s+of\s+quantities|technical\s+requirements?)\b", re.I)
IS_CITATION = re.compile(r"\bIS\s*[:\s\-]?\s*\d{2,5}\b", re.I)

FULL_TENDER_PAGES = 4     # beyond this, treat the document as a whole tender rather than a spec sheet
NAME_ONLY_WORDS = 6       # a handful of words is a product name, not a specification


class InputHandler:
    # ---------------------------------------------------------------- public

    def read(self, path: str) -> InputPayload:
        """Read any supported file into one text payload."""
        p = Path(path)
        if not p.exists():
            raise FileNotFoundError(p)
        suffix = p.suffix.lower()
        if suffix in PDF_SUFFIXES:
            return self._read_pdf(p)
        if suffix in IMAGE_SUFFIXES:
            return self._read_image(p)
        if suffix in DOC_SUFFIXES:
            return self._read_docx(p)
        if suffix in TEXT_SUFFIXES:
            return self.read_text(p.read_text(encoding="utf-8", errors="replace"), filename=p.name)
        raise ValueError(f"Unsupported input type '{suffix}'. Supported: PDF, DOCX, image, text.")

    def read_text(self, text: str, filename: Optional[str] = None,
                  source: str = "text") -> InputPayload:
        """Pasted text or a typed product name."""
        cleaned = self._clean(text)
        words = len(cleaned.split())
        return InputPayload(
            source="product_name" if words <= NAME_ONLY_WORDS and source == "text" else source,
            text=cleaned,
            pages=0,
            richness="name_only" if words <= NAME_ONLY_WORDS else "spec_only",
            filename=filename,
        )

    # ---------------------------------------------------------------- readers

    def _read_pdf(self, path: Path) -> InputPayload:
        import pymupdf

        notes: List[str] = []
        doc = pymupdf.open(str(path))
        pages = [page.get_text() for page in doc]
        doc.close()

        empty = sum(1 for p in pages if len(p.strip()) < 40)
        if empty and empty == len(pages):
            notes.append("This PDF has no text layer; it is probably a scan. OCR is needed to read it.")
        elif empty:
            notes.append(f"{empty} of {len(pages)} pages carry no text and were probably scanned.")

        text = self._clean("\n".join(pages))
        return InputPayload(
            source="pdf",
            text=text,
            pages=len(pages),
            spec_section_page=self._spec_page(pages),
            richness="full_tender" if len(pages) > FULL_TENDER_PAGES else "spec_only",
            filename=path.name,
            notes=notes,
        )

    def _read_docx(self, path: Path) -> InputPayload:
        try:
            import docx  # python-docx
        except ImportError:
            return InputPayload(source="docx", text="", richness="spec_only", filename=path.name,
                                notes=["python-docx is not installed, so this file could not be read."])
        document = docx.Document(str(path))
        blocks = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                blocks.append(" | ".join(cell.text for cell in row.cells))
        text = self._clean("\n".join(blocks))
        return InputPayload(source="docx", text=text, pages=0,
                            richness="full_tender" if len(text) > 4000 else "spec_only",
                            filename=path.name)

    def _read_image(self, path: Path) -> InputPayload:
        """A screenshot of a specification. The extracted text is meant to be shown back to the
        user for correction before anything is acted on, which is a cheap safeguard against OCR
        errors and looks honest in a demo."""
        try:
            import pytesseract
            from PIL import Image
        except ImportError:
            return InputPayload(source="image", text="", richness="spec_only", filename=path.name,
                                notes=["OCR is unavailable: install pytesseract and Tesseract to read screenshots."])
        try:
            text = pytesseract.image_to_string(Image.open(path))
        except Exception as e:      # Tesseract present as a package but not installed as a binary
            return InputPayload(source="image", text="", richness="spec_only", filename=path.name,
                                notes=[f"OCR failed: {type(e).__name__}. Check that the Tesseract binary is installed."])
        cleaned = self._clean(text)
        return InputPayload(source="image", text=cleaned, pages=1,
                            richness="name_only" if len(cleaned.split()) <= NAME_ONLY_WORDS else "spec_only",
                            filename=path.name,
                            notes=["Text was read from an image; show it back to the user before acting on it."])

    # ---------------------------------------------------------------- helpers

    @staticmethod
    def _clean(text: str) -> str:
        text = (text or "").replace("\r\n", "\n").replace(" ", " ")
        text = re.sub(r"[ \t]+", " ", text)
        text = re.sub(r"\n{3,}", "\n\n", text)
        return text.strip()

    @staticmethod
    def _spec_page(pages: List[str]) -> Optional[int]:
        """Where the technical specification starts, by heading first and IS citations second."""
        for i, page in enumerate(pages, 1):
            if SPEC_HEADING.search(page):
                return i
        best, best_count = None, 0
        for i, page in enumerate(pages, 1):
            count = len(IS_CITATION.findall(page))
            if count > best_count:
                best, best_count = i, count
        return best
